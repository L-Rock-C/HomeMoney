import sqlite3
import openpyxl
import datetime
import uuid
import re
import os

db_path = os.path.join(os.path.dirname(__file__), 'dev.db')
xlsx_path = os.path.abspath(os.path.join(os.path.dirname(__file__), '../../../data/Money.xlsx'))

print(f"Connecting to database: {db_path}")
print(f"Reading spreadsheet: {xlsx_path}")

conn = sqlite3.connect(db_path)
c = conn.cursor()

wb = openpyxl.load_workbook(xlsx_path, data_only=True)
ws_trans = [wb[s] for s in wb.sheetnames if 'Transa' in s][0]
ws_classes = [wb[s] for s in wb.sheetnames if 'Classes' in s][0]

# Clear existing data
c.execute('DELETE FROM "Transaction"')
c.execute('DELETE FROM Person')
c.execute('DELETE FROM Account')
c.execute('DELETE FROM Category')

# 1. Accounts
accounts = {'Nubank': 'CHECKING', 'Dinheiro': 'WALLET'}
acc_ids = {}
for acc, t in accounts.items():
    aid = str(uuid.uuid4())
    now_ms = int(datetime.datetime.now().timestamp() * 1000)
    c.execute('INSERT OR IGNORE INTO Account (id, name, type, color, initialBalance, createdAt, updatedAt) VALUES (?, ?, ?, ?, ?, ?, ?)',
              (aid, acc, t, '#8A05BE' if acc == 'Nubank' else '#10B981', 0, now_ms, now_ms))
    c.execute('SELECT id FROM Account WHERE name = ?', (acc,))
    acc_ids[acc] = c.fetchone()[0]

# 2. Categories
cat_rows = list(ws_classes.iter_rows(values_only=True))[1:]
cat_ids = {}
for r in cat_rows:
    cat_name = r[1]
    if cat_name:
        cid = str(uuid.uuid4())
        now_ms = int(datetime.datetime.now().timestamp() * 1000)
        is_inc = (str(cat_name).strip().lower() == 'bolsa')
        c.execute('INSERT OR IGNORE INTO Category (id, name, isIncome, createdAt, updatedAt) VALUES (?, ?, ?, ?, ?)',
                  (cid, str(cat_name).strip(), 1 if is_inc else 0, now_ms, now_ms))
        c.execute('SELECT id FROM Category WHERE name = ?', (str(cat_name).strip(),))
        cat_ids[str(cat_name).strip()] = c.fetchone()[0]

# Default category if missing
now_ms = int(datetime.datetime.now().timestamp() * 1000)
c.execute('INSERT OR IGNORE INTO Category (id, name, isIncome, createdAt, updatedAt) VALUES (?, ?, ?, ?, ?)',
          (str(uuid.uuid4()), 'Geral', 0, now_ms, now_ms))
c.execute('SELECT id FROM Category WHERE name = ?', ('Geral',))
default_cat_id = c.fetchone()[0]

# 3. People (Third-party contacts - Lucas is the User/Owner, not a contact!)
people = ['Pai', 'Lê', 'Clara', 'Matheus', 'Alfredo', 'Néfi', 'Dudu', 'Guilherme', 'Daniel', 'Hudson', 'Cactus']
person_ids = {}
for p in people:
    pid = str(uuid.uuid4())
    now_ms = int(datetime.datetime.now().timestamp() * 1000)
    c.execute('INSERT OR IGNORE INTO Person (id, name, createdAt, updatedAt) VALUES (?, ?, ?, ?)',
              (pid, p, now_ms, now_ms))
    c.execute('SELECT id FROM Person WHERE name = ?', (p,))
    person_ids[p] = c.fetchone()[0]

# 4. Transactions
rows = list(ws_trans.iter_rows(values_only=True))
print(f"Total rows in spreadsheet: {len(rows)}")
count = 0
re_inst = re.compile(r'(\d+)\s*(?:de|\/)\s*(\d+)', re.IGNORECASE)

# Clear existing transactions before seeding
c.execute('DELETE FROM "Transaction"')

for r in rows[1:]:
    if not any(r):
        continue
    submit_time = r[0]
    col_compra = r[1]
    vencimento = r[2]
    tipo_str = str(r[3]) if r[3] else 'Saída'
    nome = str(r[4]) if r[4] else 'Sem descrição'
    try:
        valor = float(r[5]) if r[5] is not None else 0.0
    except Exception:
        valor = 0.0
    cat_str = str(r[6]).strip() if r[6] else 'Geral'
    origem = str(r[7]).strip() if r[7] else ''
    destino = str(r[8]).strip() if r[8] else ''
    metodo = str(r[9]).strip() if r[9] else 'Crédito'
    pago = bool(r[10]) if r[10] is not None else True
    nota = str(r[11]) if r[11] else None

    # Determine purchase date
    p_date = col_compra if isinstance(col_compra, datetime.datetime) else (submit_time if isinstance(submit_time, datetime.datetime) else vencimento)
    if not isinstance(p_date, datetime.datetime):
        p_date = datetime.datetime.now()
    if not isinstance(vencimento, datetime.datetime):
        vencimento = p_date

    p_date_ms = int(p_date.timestamp() * 1000)
    vencimento_ms = int(vencimento.timestamp() * 1000)

    # Determine transaction type
    t_type = 'INCOME' if 'entrada' in tipo_str.lower() else 'EXPENSE'
    
    # Identify third party contact
    p_id = None
    for p in people:
        if p.lower() in origem.lower() or p.lower() in destino.lower():
            p_id = person_ids[p]
            break

    # If an expense is associated with a third party (e.g., Pai or Lelê), it's a reimbursable expense
    if t_type == 'EXPENSE' and p_id:
        if 'pai' in origem.lower() or ('lucas' in destino.lower() and 'pai' in origem.lower()):
            t_type = 'REIMBURSEMENT'
        elif 'lelê' in destino.lower() or 'lele' in destino.lower():
            t_type = 'REIMBURSEMENT'

    # Determine account
    acc_id = acc_ids['Nubank']
    if 'dinheiro' in metodo.lower():
        acc_id = acc_ids['Dinheiro']

    # Determine category
    c_id = cat_ids.get(cat_str, default_cat_id)

    # Determine installments
    curr_inst = None
    tot_inst = 1
    m = re_inst.search(nome)
    if m:
        curr_inst = int(m.group(1))
        tot_inst = int(m.group(2))

    # Payment method
    pay_m = 'CREDIT'
    if 'pix' in metodo.lower():
        pay_m = 'PIX'
    elif 'débito' in metodo.lower() or 'debito' in metodo.lower():
        pay_m = 'DEBIT'
    elif 'dinheiro' in metodo.lower():
        pay_m = 'CASH'

    tid = str(uuid.uuid4())
    now_ms = int(datetime.datetime.now().timestamp() * 1000)
    c.execute('''
        INSERT INTO "Transaction" 
        (id, description, amount, type, purchaseDate, dueDate, paid, paymentMethod, notes, currentInstallment, totalInstallments, accountId, categoryId, personId, createdAt, updatedAt)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    ''', (tid, nome, valor, t_type, p_date_ms, vencimento_ms, 1 if pago else 0, pay_m, nota, curr_inst, tot_inst, acc_id, c_id, p_id, now_ms, now_ms))
    count += 1

conn.commit()
print(f"Successfully imported {count} transactions!")
c.execute('SELECT COUNT(*) FROM "Transaction"')
print(f"Total transactions verified in DB: {c.fetchone()[0]}")
conn.close()
